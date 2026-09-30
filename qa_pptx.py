"""Editable PPTX structure, edge-case and size tests."""

import io
import unittest
import zipfile

import pandas as pd
from PIL import Image
from pptx import Presentation

from core.presentation_builder import generate_pptx


def sample_kpis(revenue_only=False):
    return {
        "gesamt_umsatz": 320000,
        "gesamt_gewinn": 64000,
        "aktuelle_marge": 20,
        "anzahl_zeilen": 12000,
        "anzahl_kategorien": 3,
        "profit_available": not revenue_only,
        "financial_aggregation_available": True,
        "kategorien_daten": pd.DataFrame([
            {"Kategorie_Clean": "Enterprise", "Umsatz_Clean": 180000, "Gewinn_Clean": 50000, "Marge": 27.8},
            {"Kategorie_Clean": "Growth", "Umsatz_Clean": 90000, "Gewinn_Clean": 12000, "Marge": 13.3},
            {"Kategorie_Clean": "Starter", "Umsatz_Clean": 50000, "Gewinn_Clean": 2000, "Marge": 4},
        ]),
        "time_analysis": {"available": True, "series": pd.DataFrame([
            {"Zeitraum": "2026-01", "Umsatz_Clean": 90000, "Gewinn_Clean": 18000},
            {"Zeitraum": "2026-02", "Umsatz_Clean": 105000, "Gewinn_Clean": 21000},
            {"Zeitraum": "2026-03", "Umsatz_Clean": 125000, "Gewinn_Clean": 25000},
        ])},
    }


class PptxTests(unittest.TestCase):
    def test_editable_deck_has_native_charts_and_tables(self):
        content = generate_pptx(sample_kpis(), {
            "zusammenfassung": "Umsatz und Ergebnis wachsen.",
            "action_plan": ["Enterprise-Angebot priorisieren", "Starter-Marge prüfen"],
        }, client_name="Nordstern GmbH", period_label="Q1 2026")
        self.assertLess(len(content), 15 * 1024 * 1024)
        prs = Presentation(io.BytesIO(content))
        self.assertEqual(len(prs.slides), 8)
        text = " ".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("Erstellt für: Nordstern GmbH", text)
        self.assertIn("Methodik und Datenqualität", text)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            names = archive.namelist()
            self.assertTrue(any(name.startswith("ppt/charts/chart") for name in names))
            self.assertFalse(any(name.endswith("vbaProject.bin") for name in names))

    def test_keynote_compatible_font_is_used(self):
        content = generate_pptx(sample_kpis(), None)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            slide_xml = b"".join(
                archive.read(name) for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )
        self.assertNotIn(b"Aptos", slide_xml)
        self.assertIn(b"Arial", slide_xml)

    def test_revenue_only_is_labelled_without_fake_margin(self):
        content = generate_pptx(sample_kpis(True), None, revenue_only=True)
        text = " ".join(shape.text for slide in Presentation(io.BytesIO(content)).slides
                        for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("Profitabilität ist für diesen Datenstand nicht berechenbar", text)

    def test_missing_series_and_negative_profit_are_supported(self):
        kpis = sample_kpis(); kpis["time_analysis"] = {}; kpis["gesamt_gewinn"] = -12000; kpis["aktuelle_marge"] = -4
        content = generate_pptx(kpis, None)
        self.assertGreater(len(content), 10_000)

    def test_safe_logo_can_be_embedded(self):
        image = Image.new("RGB", (120, 40), "white")
        buffer = io.BytesIO(); image.save(buffer, format="PNG")
        content = generate_pptx(sample_kpis(), None, logo_bytes=buffer.getvalue())
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            self.assertTrue(any(name.startswith("ppt/media/") for name in archive.namelist()))

    def test_white_label_footer_and_long_texts(self):
        content = generate_pptx(
            sample_kpis(), {"zusammenfassung": "Bericht " * 140,
                            "action_plan": ["Maßnahme " * 75]},
            client_name="Nordstern GmbH " * 8, period_label="Q1 2026",
            report_settings={"company_name": "Beratung Nord", "accent_color": "#008E78",
                             "footer_text": "Vertraulich · Beratung Nord"},
        )
        prs = Presentation(io.BytesIO(content))
        self.assertEqual(len(prs.slides), 8)
        texts = [shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text")]
        self.assertTrue(any("Vertraulich · Beratung Nord" in value for value in texts))
        self.assertTrue(any("Beratung Nord" in value for value in texts))
        self.assertTrue(any("Erstellt für: Nordstern GmbH" in value for value in texts))

    def test_consulting_polish_for_summary_numbers_and_chart_labels(self):
        kpis = sample_kpis()
        kpis["gesamt_umsatz"] = 1_312_001.96
        kpis["gesamt_gewinn"] = 778_321.27
        kpis["aktuelle_marge"] = 59.3
        kpis["anzahl_zeilen"] = 2400
        kpis["time_analysis"] = {"available": True, "comparison_available": True, "revenue_change_pct": 77.4,
                                 "margin_change_pp": -6.1, "series": pd.DataFrame([
            {"Zeitraum": pd.Timestamp("2025-10-01"), "Umsatz_Clean": 90_000, "Gewinn_Clean": 38_000},
            {"Zeitraum": pd.Timestamp("2025-11-01"), "Umsatz_Clean": 105_000, "Gewinn_Clean": 44_000},
            {"Zeitraum": pd.Timestamp("2025-12-01"), "Umsatz_Clean": 125_000, "Gewinn_Clean": 52_000},
        ])}
        content = generate_pptx(kpis, {
            "zusammenfassung": "Dieser sehr lange Absatz sollte nicht als Block auf der Management-Summary stehen. " * 12,
            "action_plan": ["Analyse der Ursachen für Margenrückgang"],
        })
        text = " ".join(shape.text for slide in Presentation(io.BytesIO(content)).slides
                        for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("Kernpunkte", text)
        self.assertIn("1,31 Mio. €", text)
        self.assertIn("778 Tsd. €", text)
        self.assertNotIn("2025-10-01 00:00:00", text)
        self.assertIn("Umsatz nach Segment in Tsd. €", text)
        self.assertIn("Margenrückgang prüfen: Kostenanstieg, Rabatte oder schwächere Nachfrage", text)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            chart_xml = b"".join(
                archive.read(name) for name in archive.namelist()
                if name.startswith("ppt/charts/chart") and name.endswith(".xml")
            ).decode("utf-8")
        self.assertIn("Okt 25", chart_xml)
        self.assertNotIn("2025-10-01 00:00:00", chart_xml)

    def test_english_report_language_translates_deck_and_formats_numbers(self):
        kpis = sample_kpis()
        kpis["gesamt_umsatz"] = 1_312_001.96
        kpis["gesamt_gewinn"] = 778_321.27
        kpis["aktuelle_marge"] = 59.3
        content = generate_pptx(
            kpis,
            {
                "zusammenfassung": "Revenue and profit are improving.",
                "action_plan": ["Review margin decline causes"],
            },
            client_name="Nordstern GmbH",
            period_label="Q1 2026",
            report_settings={"language": "en"},
        )
        text = " ".join(shape.text for slide in Presentation(io.BytesIO(content)).slides
                        for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("Prepared for: Nordstern GmbH", text)
        self.assertIn("Core metrics", text)
        self.assertIn("Revenue development", text)
        self.assertIn("Methodology and Data Quality", text)
        self.assertIn("€1.31M", text)
        self.assertIn("59.3 %", text)
        self.assertIn("Review margin decline: assess cost increases", text)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            chart_xml = b"".join(
                archive.read(name) for name in archive.namelist()
                if name.startswith("ppt/charts/chart") and name.endswith(".xml")
            ).decode("utf-8")
        self.assertIn("Jan 26", chart_xml)

    def test_short_deck_skips_detail_trend_slides(self):
        content = generate_pptx(
            sample_kpis(),
            {"action_plan": ["Management soll nächste Schritte priorisieren"]},
            report_settings={"ppt_deck_style": "short"},
        )
        prs = Presentation(io.BytesIO(content))
        self.assertEqual(len(prs.slides), 6)
        text = " ".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("Management Summary", text)
        self.assertIn("Segmentübersicht", text)
        self.assertIn("Handlungsoptionen", text)
        self.assertNotIn("Umsatzentwicklung", text)
        self.assertNotIn("Profitabilität", text)

    def test_finance_focus_and_speaker_notes_are_written(self):
        content = generate_pptx(
            sample_kpis(),
            None,
            report_settings={"ppt_audience": "finance", "ppt_speaker_notes": True},
        )
        prs = Presentation(io.BytesIO(content))
        text = " ".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("Finanz-Fokus", text)
        self.assertIn("Sprechernotizen", text)
        self.assertIn("deterministisch", text)
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            self.assertFalse(any(name.startswith("ppt/notesSlides/") for name in archive.namelist()))


if __name__ == "__main__":
    unittest.main(verbosity=2)
