"""Editable PowerPoint export built from DataDeck's deterministic KPI result."""

from __future__ import annotations

import datetime as dt
import io
import math
import re
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from .formatting import format_de_number
from .report_config import normalize_report_settings, validated_logo_bytes
from .report_language import (
    format_report_money,
    format_report_money_compact,
    format_report_number,
    format_report_percent,
    labels_for,
    normalize_report_language,
)
from .runtime_security import APP_VERSION


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
INK = RGBColor(24, 24, 33)
MUTED = RGBColor(102, 102, 119)
PAPER = RGBColor(250, 250, 252)
LINE = RGBColor(225, 225, 234)
GREEN = RGBColor(25, 160, 104)
RED = RGBColor(205, 65, 72)
PPTX_FONT = "Arial"


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def _number(value) -> float | None:
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return None


def _money(value, language: str = "de") -> str:
    number = _number(value)
    if number is None:
        return "Not available" if normalize_report_language(language) == "en" else "Nicht verfügbar"
    return format_report_money(number, language, decimals=0)


def _money_compact(value, language: str = "de") -> str:
    number = _number(value)
    if number is None:
        return "Not available" if normalize_report_language(language) == "en" else "Nicht verfügbar"
    return format_report_money_compact(number, language)


def _percent(value, language: str = "de") -> str:
    number = _number(value)
    if number is None:
        return "Not available" if normalize_report_language(language) == "en" else "Nicht verfügbar"
    return format_report_percent(number, 1, language)


def _text(slide, x, y, w, h, text: str, *, size=18, color=INK, bold=False,
          align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = str(text)
    paragraph.alignment = align
    paragraph.font.name = PPTX_FONT
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def _base_slide(prs: Presentation, title: str, accent: RGBColor, section: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(255, 255, 255)
    _text(slide, Inches(.55), Inches(.25), Inches(2.8), Inches(.25), section.upper(),
          size=9, color=accent, bold=True)
    _text(slide, Inches(.55), Inches(.55), Inches(12.1), Inches(.55), title, size=27, bold=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.55), Inches(1.15), Inches(12.2), Inches(.025))
    line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background()
    return slide


def _footer(slide, page: int, brand: str, period: str, footer_text: str = ""):
    context = footer_text or period or "Analysezeitraum nicht angegeben"
    _text(slide, Inches(.55), Inches(7.12), Inches(7), Inches(.18),
          f"{brand[:35]}  |  {context[:65]}", size=8, color=MUTED)
    _text(slide, Inches(12.1), Inches(7.12), Inches(.65), Inches(.18), str(page),
          size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def _speaker_notes(notes_bucket: list[str], notes: str, enabled: bool) -> None:
    if enabled and notes:
        notes_bucket.append(_clean_slide_sentence(notes)[:260])


def _append_speaker_notes_slide(
    prs: Presentation,
    notes: list[str],
    accent: RGBColor,
    labels: dict,
    language: str,
    brand: str,
    period_label: str,
    footer_text: str,
    page: int,
) -> None:
    if not notes:
        return
    title = "Presenter notes" if language == "en" else "Sprechernotizen"
    section = "Appendix" if language == "en" else "Anhang"
    slide = _base_slide(prs, title, accent, section)
    intro = (
        "Use these prompts while presenting the deck. They are kept as an editable slide for Keynote compatibility."
        if language == "en" else
        "Diese Hinweise helfen beim Vorstellen des Decks. Sie bleiben als editierbare Folie erhalten, damit Keynote den Export sauber importiert."
    )
    _text(slide, Inches(.75), Inches(1.45), Inches(11.4), Inches(.55), intro, size=15, color=MUTED)
    _bullets(slide, Inches(.75), Inches(2.15), Inches(11.4), Inches(4.45), notes[:7], accent, 14)
    _footer(slide, page, brand, period_label, footer_text)


def _card(slide, x, y, w, h, label: str, value: str, accent: RGBColor, note: str = ""):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = PAPER
    shape.line.color.rgb = LINE
    shape.adjustments[0] = 0.08
    _text(slide, x + Inches(.2), y + Inches(.18), w - Inches(.4), Inches(.2),
          label.upper(), size=8, color=MUTED, bold=True)
    _text(slide, x + Inches(.2), y + Inches(.53), w - Inches(.4), Inches(.45),
          value, size=22, color=INK, bold=True)
    if note:
        _text(slide, x + Inches(.2), y + h - Inches(.38), w - Inches(.4), Inches(.2),
              note, size=8, color=accent)


def _bullets(slide, x, y, w, h, items: list[str], accent: RGBColor, size: int = 15):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True
    for index, item in enumerate(items[:6]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = _clean_slide_sentence(str(item))[:260]
        paragraph.font.name = PPTX_FONT; paragraph.font.size = Pt(size); paragraph.font.color.rgb = INK
        paragraph.space_after = Pt(12); paragraph.level = 0
        paragraph.text = f"•  {paragraph.text}"
    return box


def _chart_heading(slide, x, y, w, title: str, unit: str | None, accent: RGBColor):
    heading = title if not unit else f"{title} in {unit}"
    _text(slide, x, y, w, Inches(.25), heading, size=12, color=accent, bold=True)


def _chart(slide, categories: list[str], values: list[float], x, y, w, h,
           accent: RGBColor, title: str, chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED,
           no_data: str = "Keine vergleichbaren Daten verfügbar."):
    if not categories or not values:
        _text(slide, x, y + Inches(.7), w, Inches(.5), no_data,
              size=14, color=MUTED, align=PP_ALIGN.CENTER)
        return None
    data = ChartData(); data.categories = categories
    data.add_series(title, values)
    chart = slide.shapes.add_chart(chart_type, x, y, w, h, data).chart
    chart.has_legend = False
    chart.has_title = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = LINE
    chart.value_axis.tick_labels.font.name = PPTX_FONT; chart.value_axis.tick_labels.font.size = Pt(9)
    chart.category_axis.tick_labels.font.name = PPTX_FONT; chart.category_axis.tick_labels.font.size = Pt(9)
    chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb = accent
    chart.series[0].format.line.color.rgb = accent
    return chart


def _table(slide, rows: list[tuple[str, str, str]], x, y, w, h, accent: RGBColor, labels: dict):
    table = slide.shapes.add_table(max(2, len(rows) + 1), 3, x, y, w, h).table
    table.columns[0].width = int(w * .48); table.columns[1].width = int(w * .29); table.columns[2].width = int(w * .23)
    for col, text in enumerate((labels["segment"], labels["revenue"], labels["margin"])):
        cell = table.cell(0, col); cell.text = text
        cell.fill.solid(); cell.fill.fore_color.rgb = accent
        for run in cell.text_frame.paragraphs[0].runs:
            run.font.name = PPTX_FONT; run.font.size = Pt(10); run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    for row_index, row in enumerate(rows[:6], start=1):
        for col, text in enumerate(row):
            cell = table.cell(row_index, col); cell.text = text
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(255,255,255) if row_index % 2 else PAPER
            for run in cell.text_frame.paragraphs[0].runs:
                run.font.name = PPTX_FONT; run.font.size = Pt(10); run.font.color.rgb = INK
    return table


def _period_label(value: Any, language: str = "de") -> str:
    if value in (None, ""):
        return ""
    try:
        timestamp = pd.to_datetime(value, errors="coerce")
    except (TypeError, ValueError):
        timestamp = pd.NaT
    if pd.notna(timestamp):
        months = (
            ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
            if normalize_report_language(language) == "en"
            else ("Jan", "Feb", "Mrz", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez")
        )
        return f"{months[timestamp.month - 1]} {str(timestamp.year)[-2:]}"
    return str(value).split(" ")[0][:10]


def _series_from_kpis(kpis: dict, language: str = "de") -> tuple[list[str], list[float], list[float]]:
    frame = (kpis.get("time_analysis") or {}).get("series")
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return [], [], []
    rows = frame.tail(12).to_dict("records")
    categories = [_period_label(row.get("Zeitraum", ""), language) for row in rows]
    revenues = [_number(row.get("Umsatz_Clean")) or 0 for row in rows]
    profits = [_number(row.get("Gewinn_Clean")) or 0 for row in rows]
    return categories, revenues, profits


def _segments_from_kpis(kpis: dict) -> list[dict[str, Any]]:
    frame = kpis.get("kategorien_daten")
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return []
    return frame.sort_values("Umsatz_Clean", ascending=False).head(5).to_dict("records")


def _money_axis(values: list[float], language: str = "de") -> tuple[list[float], str]:
    if not values:
        return values, "€"
    maximum = max(abs(value) for value in values)
    if normalize_report_language(language) == "en":
        if maximum >= 1_000_000:
            return [value / 1_000_000 for value in values], "€M"
        if maximum >= 10_000:
            return [value / 1_000 for value in values], "€K"
        return values, "€"
    if maximum >= 1_000_000:
        return [value / 1_000_000 for value in values], "Mio. €"
    if maximum >= 10_000:
        return [value / 1_000 for value in values], "Tsd. €"
    return values, "€"


def _clean_slide_sentence(text: str) -> str:
    cleaned = re.sub(r"\s+", " ", str(text or "")).strip()
    cleaned = cleaned.replace("**", "").replace("BEOBACHTUNG", "").replace("BEWERTUNG", "")
    cleaned = cleaned.replace("EMPFEHLUNG", "").strip(" :-")
    return cleaned


def _polish_action(text: str, language: str = "de") -> str:
    cleaned = _clean_slide_sentence(text)
    lower = cleaned.lower()
    if normalize_report_language(language) == "en":
        if ("margin decline" in lower or "margin drop" in lower) and ("cause" in lower or "analy" in lower):
            return "Review margin decline: assess cost increases, discounting or weaker demand as hypotheses."
        if "causes" in lower and "review" not in lower:
            return f"{cleaned}: review the responsible KPI, period and segment."
        return cleaned
    if "margenrückgang" in lower and ("ursache" in lower or "analyse" in lower):
        return "Margenrückgang prüfen: Kostenanstieg, Rabatte oder schwächere Nachfrage als Hypothesen bewerten."
    if "ursachen" in lower and "prüfen" not in lower:
        return f"{cleaned}: verantwortliche Kennzahl, Zeitraum und Segment gezielt prüfen."
    return cleaned


def _action_items(ai_insights: dict | None, consultant_comment: str, language: str = "de") -> list[str]:
    actions = [_polish_action(item, language) for item in list((ai_insights or {}).get("action_plan") or [])]
    if consultant_comment.strip():
        actions.insert(0, _clean_slide_sentence(consultant_comment))
    return [item for item in actions if item][:5]


def _audience_priority_title(audience: str, labels: dict, language: str) -> str:
    if language == "en":
        return {
            "finance": "Financial focus",
            "growth": "Growth focus",
            "management": "Priorities",
        }.get(audience, labels["priorities"])
    return {
        "finance": "Finanz-Fokus",
        "growth": "Wachstums-Fokus",
        "management": "Prioritäten",
    }.get(audience, labels["priorities"])


def _audience_fallback_actions(audience: str, language: str) -> list[str]:
    if language == "en":
        return {
            "finance": [
                "Validate margin movement against cost base and accounting logic",
                "Review outliers before using the report externally",
            ],
            "growth": [
                "Prioritize the strongest revenue segments for the next client discussion",
                "Review weak segments for offer, pricing or demand issues",
            ],
            "management": [
                "Review KPIs with the business owner",
                "Assign owners and timing to next actions",
            ],
        }.get(audience, [])
    return {
        "finance": [
            "Margenentwicklung mit Kostenbasis und Buchungslogik validieren",
            "Ausreißer vor externer Verwendung fachlich prüfen",
        ],
        "growth": [
            "Umsatzstärkste Segmente für das nächste Kundengespräch priorisieren",
            "Schwache Segmente auf Angebot, Preis oder Nachfrage prüfen",
        ],
        "management": [
            "Kennzahlen fachlich prüfen",
            "Maßnahmen mit Verantwortlichen und Termin versehen",
        ],
    }.get(audience, [])


def _first_performer_name(kpis: dict, segments: list[dict[str, Any]]) -> str | None:
    performers = kpis.get("top_performer") or []
    if performers:
        return str(performers[0].get("Kategorie_Clean") or "")[:40]
    if segments:
        return str(segments[0].get("Kategorie_Clean") or "")[:40]
    return None


def _summary_points(
    kpis: dict,
    ai_insights: dict | None,
    revenue_only: bool,
    aggregation: bool,
    segments: list[dict[str, Any]],
    language: str = "de",
) -> list[str]:
    lang = normalize_report_language(language)
    points: list[str] = []
    if aggregation:
        points.append(
            f"Revenue is {_money_compact(kpis.get('gesamt_umsatz'), lang)}"
            if lang == "en" else
            f"Umsatz liegt bei {_money_compact(kpis.get('gesamt_umsatz'), lang)}"
        )
        if not revenue_only:
            points.append(
                f"Profit is {_money_compact(kpis.get('gesamt_gewinn'), lang)}, margin is {_percent(kpis.get('aktuelle_marge'), lang)}"
                if lang == "en" else
                f"Gewinn liegt bei {_money_compact(kpis.get('gesamt_gewinn'), lang)}, Marge bei {_percent(kpis.get('aktuelle_marge'), lang)}"
            )
    else:
        points.append(
            "Multiple currencies detected, financial totals are not aggregated"
            if lang == "en" else
            "Mehrere Währungen erkannt, Finanzsummen werden nicht aggregiert"
        )

    time_analysis = kpis.get("time_analysis") or {}
    if time_analysis.get("comparison_available"):
        points.append(
            f"Revenue in the comparison period {_percent(time_analysis.get('revenue_change_pct'), lang)}"
            if lang == "en" else
            f"Umsatz im Vergleichszeitraum {_percent(time_analysis.get('revenue_change_pct'), lang)}"
        )
        if not revenue_only and _number(time_analysis.get("margin_change_pp")) is not None:
            points.append(
                f"Margin change {format_report_number(time_analysis.get('margin_change_pp'), 1, lang)} percentage points"
                if lang == "en" else
                f"Margenveränderung {format_de_number(time_analysis.get('margin_change_pp'), 1)} Prozentpunkte"
            )

    top_name = _first_performer_name(kpis, segments)
    if top_name:
        if lang == "en":
            metric = "revenue driver" if revenue_only else "profit driver"
        else:
            metric = "Umsatztreiber" if revenue_only else "Gewinntreiber"
        points.append(
            f"{top_name} is the strongest {metric}"
            if lang == "en" else
            f"{top_name} ist der stärkste {metric}"
        )

    summary = _clean_slide_sentence((ai_insights or {}).get("zusammenfassung") or "")
    if len(points) < 4 and summary:
        first_sentence = re.split(r"(?<=[.!?])\s+", summary)[0]
        if first_sentence:
            points.append(first_sentence[:180])
    return points[:4]


def generate_pptx(
    kpis: dict,
    ai_insights: dict | None,
    niche: str | None = None,
    revenue_only: bool = False,
    data_quality: dict | None = None,
    column_mapping: dict | None = None,
    report_settings: dict | None = None,
    consultant_comment: str = "",
    client_name: str = "",
    period_label: str = "",
    report_version: int = 1,
    logo_bytes: bytes | None = None,
) -> bytes:
    """Build an editable consulting deck without recomputing financial KPIs."""
    settings = normalize_report_settings(report_settings)
    language = normalize_report_language(settings.get("language"))
    labels = labels_for(language)
    deck_style = settings.get("ppt_deck_style", "full")
    short_deck = deck_style == "short"
    audience = settings.get("ppt_audience", "management")
    speaker_notes = bool(settings.get("ppt_speaker_notes", False))
    accent = _rgb(settings["accent_color"])
    brand = settings["company_name"] or "DataDeck"
    footer_text = settings["footer_text"]
    logo = validated_logo_bytes(logo_bytes)
    prs = Presentation(); prs.slide_width = SLIDE_W; prs.slide_height = SLIDE_H
    presenter_notes: list[str] = []

    # 1 Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = INK
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(.75), Inches(.12), Inches(5.9))
    band.fill.solid(); band.fill.fore_color.rgb = accent; band.line.fill.background()
    if logo:
        slide.shapes.add_picture(io.BytesIO(logo), Inches(10.6), Inches(.65), width=Inches(1.8))
    _text(slide, Inches(1.05), Inches(1.35), Inches(10.6), Inches(.5), brand, size=16,
          color=RGBColor(255,255,255), bold=True)
    _text(slide, Inches(1.05), Inches(2.15), Inches(10.6), Inches(1.25),
          labels["cover_subtitle"].replace(" & ", "\n"), size=34, color=RGBColor(255,255,255), bold=True)
    client_label = f"{labels['prepared_for']}: {client_name}" if client_name else f"{labels['prepared_for']}: {labels['client_default']}"
    _text(slide, Inches(1.05), Inches(4.0), Inches(8.8), Inches(.9),
          f"{client_label}\n{period_label or labels['current_period']}", size=15,
          color=RGBColor(205,205,218))
    _text(slide, Inches(1.05), Inches(6.55), Inches(9), Inches(.25),
          f"{labels['report_version']} {max(1, int(report_version))}  |  {dt.date.today().strftime('%Y-%m-%d' if language == 'en' else '%d.%m.%Y')}",
          size=9, color=RGBColor(170,170,188))
    _speaker_notes(
        presenter_notes,
        "Open with the client, period and purpose of the analysis." if language == "en"
        else "Mit Mandant, Zeitraum und Ziel der Analyse eröffnen.",
        speaker_notes,
    )

    # 2 Executive summary
    page = 2
    slide = _base_slide(prs, labels["management_summary"], accent, labels["interpretation"])
    aggregation = bool(kpis.get("financial_aggregation_available", True))
    segments = _segments_from_kpis(kpis)
    actions = _action_items(ai_insights, consultant_comment, language)
    _text(slide, Inches(.7), Inches(1.45), Inches(7.4), Inches(.3), labels["key_points"], size=12, color=accent, bold=True)
    _bullets(
        slide,
        Inches(.7),
        Inches(1.88),
        Inches(7.2),
        Inches(4.45),
        _summary_points(kpis, ai_insights, revenue_only, aggregation, segments, language),
        accent,
        17,
    )
    _text(slide, Inches(8.65), Inches(1.45), Inches(3.75), Inches(.3), _audience_priority_title(audience, labels, language), size=12, color=accent, bold=True)
    _bullets(slide, Inches(8.65), Inches(1.88), Inches(3.75), Inches(4.45), actions or _audience_fallback_actions(audience, language), accent, 13)
    _speaker_notes(
        presenter_notes,
        "Focus this slide on the decision relevance, not on every detail." if language == "en"
        else "Diese Folie auf Entscheidungsrelevanz fokussieren, nicht auf jedes Detail.",
        speaker_notes,
    )
    _footer(slide, page, brand, period_label, footer_text); page += 1

    # 3 KPI overview
    slide = _base_slide(prs, labels["core_metrics"], accent, labels["performance"])
    _card(slide, Inches(.65), Inches(1.55), Inches(2.9), Inches(1.55), labels["revenue"],
          _money_compact(kpis.get("gesamt_umsatz"), language) if aggregation else labels["multiple_currencies"], accent)
    _card(slide, Inches(3.75), Inches(1.55), Inches(2.9), Inches(1.55), labels["profit"],
          labels["not_available"] if revenue_only or not aggregation else _money_compact(kpis.get("gesamt_gewinn"), language), accent)
    _card(slide, Inches(6.85), Inches(1.55), Inches(2.9), Inches(1.55), labels["margin"],
          labels["not_available"] if revenue_only or not aggregation else _percent(kpis.get("aktuelle_marge"), language), accent)
    _card(slide, Inches(9.95), Inches(1.55), Inches(2.7), Inches(1.55), labels["records"],
          format_report_number(int(kpis.get('anzahl_zeilen', 0)), 0, language), accent)
    note = (
        "Profit and margin cannot be calculated without a cost or profit column."
        if language == "en" and revenue_only else
        "Gewinn und Marge sind ohne Kosten- oder Gewinnspalte nicht berechenbar."
        if revenue_only else (
            "Financial totals are not aggregated across multiple currencies without documented exchange rates."
            if language == "en" and not aggregation else
            "Finanzsummen werden bei mehreren Währungen ohne dokumentierte Wechselkurse bewusst nicht aggregiert."
            if not aggregation else
            "All financial KPIs were calculated deterministically from the checked data."
            if language == "en" else
            "Alle Finanzkennzahlen wurden deterministisch aus dem geprüften Datenstand berechnet."
        )
    )
    _text(slide, Inches(.7), Inches(3.65), Inches(11.8), Inches(.7), note, size=17, color=MUTED)
    _speaker_notes(
        presenter_notes,
        "Mention that DataDeck calculated these KPIs deterministically before AI interpretation." if language == "en"
        else "Kurz erwähnen, dass DataDeck diese Kennzahlen deterministisch vor der KI-Einordnung berechnet.",
        speaker_notes,
    )
    _footer(slide, page, brand, period_label, footer_text); page += 1

    # 4 Revenue
    categories, revenues, profits = _series_from_kpis(kpis, language)
    if not short_deck:
        slide = _base_slide(prs, labels["revenue_development"], accent, labels["timeline"])
        revenue_values, revenue_unit = _money_axis(revenues, language)
        _chart_heading(slide, Inches(.65), Inches(1.38), Inches(8.4), labels["revenue_development"], revenue_unit, accent)
        _chart(slide, categories, revenue_values, Inches(.65), Inches(1.75), Inches(8.4), Inches(4.65), accent, labels["revenue"], no_data=labels["no_chart_data"])
        _text(slide, Inches(9.45), Inches(1.6), Inches(3.0), Inches(.3), labels["reading"], size=12, color=accent, bold=True)
        time_info = kpis.get("time_analysis") or {}
        commentary = time_info.get("reason") or (
            "Shown are up to twelve available analysis periods."
            if language == "en" else
            "Dargestellt werden bis zu zwölf verfügbare Analyseperioden."
        )
        _text(slide, Inches(9.45), Inches(2.05), Inches(3.0), Inches(2.5), str(commentary)[:600], size=15, color=INK)
        _speaker_notes(
            presenter_notes,
            "Use this slide to explain the trend and any period comparison caveats." if language == "en"
            else "Diese Folie für Trend und eventuelle Vergleichseinschränkungen nutzen.",
            speaker_notes,
        )
        _footer(slide, page, brand, period_label, footer_text); page += 1

        # 5 Profitability
        slide = _base_slide(prs, labels["profitability"], accent, labels["result"])
        if revenue_only:
            _text(slide, Inches(.75), Inches(2.0), Inches(11.8), Inches(1.1),
                  "Profitability cannot be calculated for this dataset." if language == "en" else "Profitabilität ist für diesen Datenstand nicht berechenbar.", size=25, bold=True)
            _text(slide, Inches(.75), Inches(3.25), Inches(10.5), Inches(.8),
                  "Profit and margin analysis requires a cost, COGS or profit column." if language == "en" else "Für Gewinn- und Margenanalysen wird eine Kosten-, COGS- oder Gewinnspalte benötigt.",
                  size=17, color=MUTED)
        else:
            profit_values, profit_unit = _money_axis(profits, language)
            _chart_heading(slide, Inches(.65), Inches(1.38), Inches(8.4), labels["profit_development"], profit_unit, GREEN)
            _chart(slide, categories, profit_values, Inches(.65), Inches(1.75), Inches(8.4), Inches(4.65), GREEN,
                   labels["profit"], no_data=labels["no_chart_data"])
            _card(slide, Inches(9.45), Inches(1.65), Inches(3.0), Inches(1.55), labels["total_margin"],
                  _percent(kpis.get("aktuelle_marge"), language), accent)
            profit = _number(kpis.get("gesamt_gewinn"))
            color = GREEN if profit is not None and profit >= 0 else RED
            _text(slide, Inches(9.55), Inches(3.6), Inches(2.8), Inches(.4),
                  labels["positive_result"] if profit is not None and profit >= 0 else labels["negative_result"],
                  size=15, color=color, bold=True)
        _speaker_notes(
            presenter_notes,
            "Frame profitability carefully when cost or profit columns are incomplete." if language == "en"
            else "Profitabilität vorsichtig einordnen, wenn Kosten- oder Gewinnspalten unvollständig sind.",
            speaker_notes,
        )
        _footer(slide, page, brand, period_label, footer_text); page += 1

    # 6 Segments
    slide = _base_slide(prs, labels["segment_overview"], accent, labels["portfolio"])
    rows = [(str(row.get("Kategorie_Clean", ""))[:42], _money_compact(row.get("Umsatz_Clean"), language),
             "-" if revenue_only else _percent(row.get("Marge"), language)) for row in segments]
    if rows:
        _table(slide, rows, Inches(.65), Inches(1.55), Inches(7.2), Inches(4.8), accent, labels)
        segment_values, segment_unit = _money_axis([_number(item.get("Umsatz_Clean")) or 0 for item in segments], language)
        _chart_heading(slide, Inches(8.15), Inches(1.42), Inches(4.5), labels["revenue_by_segment"], segment_unit, accent)
        _chart(slide, [row[0] for row in rows], segment_values,
               Inches(8.15), Inches(1.8), Inches(4.5), Inches(4.55), accent, labels["revenue"], no_data=labels["no_chart_data"])
    else:
        _text(slide, Inches(.75), Inches(2.2), Inches(11.7), Inches(.6),
              labels["no_segments"], size=22, color=MUTED, align=PP_ALIGN.CENTER)
    _speaker_notes(
        presenter_notes,
        "Use segment concentration to discuss priorities and possible follow-up questions." if language == "en"
        else "Segmentkonzentration nutzen, um Prioritäten und Folgefragen zu besprechen.",
        speaker_notes,
    )
    _footer(slide, page, brand, period_label, footer_text); page += 1

    # 7 Insights
    slide = _base_slide(prs, labels["recommendations"], accent, labels["ai_insights"])
    _bullets(slide, Inches(.75), Inches(1.55), Inches(11.7), Inches(4.95), actions or [
        "No AI interpretation has been created yet. The KPIs can still be reviewed and commented on."
        if language == "en" else
        "Noch keine KI-Einordnung erstellt. Die Kennzahlen können unabhängig davon geprüft und kommentiert werden."
    ], accent, 17)
    _text(slide, Inches(.75), Inches(6.45), Inches(11.6), Inches(.3),
          "AI content only interprets aggregated KPIs; financial KPIs are not calculated by AI."
          if language == "en" else
          "KI-Inhalte interpretieren ausschließlich aggregierte Kennzahlen; Finanz-KPIs werden nicht von der KI berechnet.",
          size=9, color=MUTED)
    _speaker_notes(
        presenter_notes,
        "Turn recommendations into owners, decisions and next steps during the meeting." if language == "en"
        else "Empfehlungen im Termin in Verantwortliche, Entscheidungen und nächste Schritte übersetzen.",
        speaker_notes,
    )
    _footer(slide, page, brand, period_label, footer_text); page += 1

    # 8 Methodology
    slide = _base_slide(prs, labels["methodology"], accent, labels["appendix"])
    quality = data_quality or {}
    if language == "en":
        quality_items = [
            f"Analysis quality: {quality.get('level', 'not assessed')}",
            f"Checked data basis: {format_report_number(int(kpis.get('anzahl_zeilen', 0)), 0, language)} records",
            "KPIs are calculated deterministically in Python.",
            labels["method_no_raw"],
            labels["disclaimer"],
        ]
    else:
        quality_items = [
            f"Analysequalität: {quality.get('level', 'nicht bewertet')}",
            f"Geprüfte Datenbasis: {int(kpis.get('anzahl_zeilen', 0)):,} Datensätze".replace(",", "."),
            "Kennzahlen werden deterministisch in Python berechnet.",
            labels["method_no_raw"],
            labels["disclaimer"],
        ]
    _bullets(slide, Inches(.75), Inches(1.5), Inches(7.5), Inches(4.8), quality_items, accent, 15)
    mapping = column_mapping or kpis.get("column_mapping") or {}
    _text(slide, Inches(8.65), Inches(1.55), Inches(3.6), Inches(.3), labels["used_columns"], size=12, color=accent, bold=True)
    mapping_labels = (("umsatz", labels["revenue"]), ("gewinn", labels["profit"]), ("kosten", labels["costs"]), ("datum", labels["date"]))
    mapping_text = "\n".join(f"{label}: {mapping.get(key) or '-'}" for key, label in mapping_labels)
    _text(slide, Inches(8.65), Inches(2.0), Inches(3.6), Inches(2.7), mapping_text, size=14)
    _text(slide, Inches(8.65), Inches(5.35), Inches(3.6), Inches(.5), f"DataDeck Engine {APP_VERSION}", size=9, color=MUTED)
    _speaker_notes(
        presenter_notes,
        "This is the trust slide: deterministic KPIs, aggregated AI context and data-quality limits." if language == "en"
        else "Das ist die Vertrauensfolie: deterministische KPIs, aggregierter KI-Kontext und Datenqualitätsgrenzen.",
        speaker_notes,
    )
    _footer(slide, page, brand, period_label, footer_text); page += 1

    _append_speaker_notes_slide(
        prs, presenter_notes, accent, labels, language, brand, period_label, footer_text, page
    )

    output = io.BytesIO(); prs.save(output)
    content = output.getvalue()
    if len(content) > 15 * 1024 * 1024:
        raise ValueError("Die PPTX-Datei überschreitet das sichere Größenlimit")
    return content
